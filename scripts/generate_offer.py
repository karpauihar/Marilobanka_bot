from __future__ import annotations

import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "public_offer.md"
DOCX_OUTPUT = ROOT / "docs" / "public_offer.docx"
PDF_OUTPUT = ROOT / "docs" / "public_offer.pdf"

LINK_RE = re.compile(r"\[([^]]+)]\(([^)]+)\)")


def plain_text(value: str) -> str:
    def replace_link(match: re.Match[str]) -> str:
        label, url = match.groups()
        return label if label == url else f"{label} ({url})"

    return LINK_RE.sub(replace_link, value)


def parse_source() -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            blocks.append(("heading", plain_text(line[4:])))
        elif line.startswith("## "):
            blocks.append(("subtitle", plain_text(line[3:])))
        elif line.startswith("# "):
            blocks.append(("title", plain_text(line[2:])))
        elif line.startswith("Редакция от "):
            blocks.append(("date", plain_text(line)))
        else:
            blocks.append(("body", plain_text(line).replace("  ", " ")))
    return blocks


def set_cell_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_cell_font(run, "Arial", 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def build_docx(blocks: list[tuple[str, str]]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    if "Offer Heading" not in document.styles:
        heading_style = document.styles.add_style(
            "Offer Heading", WD_STYLE_TYPE.PARAGRAPH
        )
    else:
        heading_style = document.styles["Offer Heading"]
    heading_style.font.name = "Arial"
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    heading_style.paragraph_format.space_before = Pt(12)
    heading_style.paragraph_format.space_after = Pt(6)
    heading_style.paragraph_format.keep_with_next = True

    for kind, text in blocks:
        if kind == "title":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            set_cell_font(paragraph.add_run(text.upper()), "Arial", 16, True)
        elif kind == "subtitle":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            set_cell_font(paragraph.add_run(text), "Arial", 12, True)
        elif kind == "date":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(16)
            run = paragraph.add_run(text)
            set_cell_font(run, "Arial", 10)
            run.italic = True
        elif kind == "heading":
            document.add_paragraph(text, style="Offer Heading")
        else:
            paragraph = document.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0.75)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    document.core_properties.title = "Публичная оферта — Домашний уход за стопами"
    document.core_properties.author = "Лобанова Мария Андреевна"
    document.save(DOCX_OUTPUT)


def find_pdf_fonts() -> tuple[Path, Path]:
    candidates = [
        (
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
    ]
    for regular, bold in candidates:
        if regular.is_file() and bold.is_file():
            return regular, bold
    raise RuntimeError("Arial or DejaVu Sans fonts are required to build the PDF")


def draw_page_number(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("OfferRegular", 9)
    canvas.drawCentredString(A4[0] / 2, 1.1 * cm, str(document.page))
    canvas.restoreState()


def build_pdf(blocks: list[tuple[str, str]]) -> None:
    regular_font, bold_font = find_pdf_fonts()
    pdfmetrics.registerFont(TTFont("OfferRegular", regular_font))
    pdfmetrics.registerFont(TTFont("OfferBold", bold_font))

    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "OfferTitle",
        parent=styles["Title"],
        fontName="OfferBold",
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    subtitle = ParagraphStyle(
        "OfferSubtitle",
        parent=styles["Heading2"],
        fontName="OfferBold",
        fontSize=12,
        leading=15,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    date = ParagraphStyle(
        "OfferDate",
        parent=styles["Normal"],
        fontName="OfferRegular",
        fontSize=10,
        leading=13,
        alignment=TA_CENTER,
        spaceAfter=16,
    )
    heading = ParagraphStyle(
        "OfferHeading",
        parent=styles["Heading3"],
        fontName="OfferBold",
        fontSize=12,
        leading=15,
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True,
    )
    body = ParagraphStyle(
        "OfferBody",
        parent=styles["BodyText"],
        fontName="OfferRegular",
        fontSize=10.2,
        leading=13.2,
        alignment=TA_JUSTIFY,
        firstLineIndent=0.75 * cm,
        spaceAfter=5,
    )

    story = []
    style_map = {
        "title": title,
        "subtitle": subtitle,
        "date": date,
        "heading": heading,
        "body": body,
    }
    for kind, text in blocks:
        story.append(Paragraph(escape(text), style_map[kind]))

    pdf = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=2.2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title="Публичная оферта — Домашний уход за стопами",
        author="Лобанова Мария Андреевна",
    )
    pdf.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)


def main() -> None:
    blocks = parse_source()
    build_docx(blocks)
    build_pdf(blocks)
    print(f"Created {DOCX_OUTPUT.relative_to(ROOT)}")
    print(f"Created {PDF_OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
